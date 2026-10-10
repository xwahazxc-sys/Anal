"""Рендер в ATS-безопасные форматы: простой Markdown и одноколоночный DOCX."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from .models import Profile, TailoredResume


def to_markdown(p: Profile, r: TailoredResume) -> str:
    c = p.contacts
    contact = " | ".join(x for x in [c.get("email"), c.get("phone"), c.get("location"), *c.get("links", [])] if x)
    out = [f"# {p.name}", r.headline, contact, "", "## Summary", r.summary, "", "## Skills", ", ".join(r.skills), "", "## Experience"]
    for e in r.experience:
        out += [f"### {e.title} — {e.company} ({e.start} – {e.end})", *[f"- {b}" for b in e.bullets], ""]
    out.append("## Education")
    out += [f"- {e.school}, {e.degree} {e.year or ''}".strip() for e in p.education]
    if p.languages:
        out += ["", "## Languages", ", ".join(p.languages)]
    return "\n".join(out) + "\n"


def to_docx(p: Profile, r: TailoredResume, path: Path) -> None:
    d = Document()
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_heading(p.name, 0)
    d.add_paragraph(r.headline)
    c = p.contacts
    d.add_paragraph(" | ".join(x for x in [c.get("email"), c.get("phone"), c.get("location"), *c.get("links", [])] if x))
    d.add_heading("Summary", 1); d.add_paragraph(r.summary)
    d.add_heading("Skills", 1); d.add_paragraph(", ".join(r.skills))
    d.add_heading("Experience", 1)
    for e in r.experience:
        d.add_heading(f"{e.title} — {e.company} ({e.start} – {e.end})", 2)
        for b in e.bullets:
            d.add_paragraph(b, style="List Bullet")
    d.add_heading("Education", 1)
    for e in p.education:
        d.add_paragraph(f"{e.school}, {e.degree} {e.year or ''}".strip())
    if p.languages:
        d.add_heading("Languages", 1); d.add_paragraph(", ".join(p.languages))
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)
